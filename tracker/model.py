"""The welfare model's calculations, in Python, for workbooks saved without calculated values.

The Wales workbook was built by a script, so its formula cells carry no saved results: openpyxl
reads them as empty until the file is opened and saved in Excel. This module recomputes them from
the workbook's own inputs, following its named functions (TaxOnGross, BandMeanTax,
ParetoBandMeanTax) and cell formulas line for line.

    python3 model.py            checks these functions against the England workbook's saved results
"""
import math
import os
import re
import statistics

WEIGHTS = [10, 10, 5, 5, 10, 20, 10, 5, 5, 10, 10]  # percentile band widths, P0-10 ... P90-100
BENEFIT_SHEETS = ["JSA", "Universal Credit", "State Pension", "Pension Credit", "Housing Benefit", "PIP",
                  "Disability Living Allowance", "Employment Support Allowance", "Carers Allowance",
                  "Attendance Allowance"]


def xround(x, nd=2):
    """Excel ROUND: half away from zero."""
    m = 10 ** nd
    return math.copysign(math.floor(abs(x) * m + 0.5) / m, x)


class Model:
    def __init__(self, p):
        self.p = p  # tax parameters, by the workbook's names (Tax_PA, Tax_BasicBand, ...)

    def tax_on_gross(self, g):
        p = self.p
        sac = max(0, g * p["Tax_SacPct"])
        e = max(0, g - sac)
        if e <= p["Tax_TaperStart"]:
            pa = p["Tax_PA"]
        elif e >= p["Tax_TaperZero"]:
            pa = 0
        else:
            pa = max(0, p["Tax_PA"] - math.floor((e - p["Tax_TaperStart"]) / 2))
        t = max(0, e - pa)
        band, add = p["Tax_BasicBand"], p["Tax_AddThreshold"]
        it = xround(min(t, band) * p["Tax_BasicRate"] + min(max(t - band, 0), add - band) * p["Tax_HigherRate"]
                    + max(0, t - band - (add - band)) * p["Tax_AddRate"], 2)
        nb = max(0, g - min(sac, p["Tax_NISacCap"]))
        ni = xround((max(min(nb, p["Tax_NIUEL"]) - p["Tax_NIPT"], 0) * p["Tax_NIMain"]
                     + max(0, nb - p["Tax_NIUEL"]) * p["Tax_NIAdd"]) if nb > p["Tax_NIPT"] else 0, 2)
        return xround(it + ni, 2)

    def band_mean_tax(self, lo, hi):
        p, s = self.p, self.p["Tax_SacPct"]
        brk = [v / (1 - s) for v in (p["Tax_PA"], p["Tax_PA"] + p["Tax_BasicBand"], p["Tax_TaperStart"],
                                     p["Tax_TaperZero"], p["Tax_AddThreshold"], p["Tax_NIPT"], p["Tax_NIUEL"])]
        brk += [p["Tax_NIPT"] + p["Tax_NISacCap"], p["Tax_NIUEL"] + p["Tax_NISacCap"],
                p["Tax_NISacCap"] / s if s > 0 else lo]
        if hi <= lo:
            return self.tax_on_gross(lo)
        grid = [lo + (hi - lo) * k / 8 for k in range(9)]
        inside = [b for b in brk if lo < b < hi] or [lo]
        pts = sorted(set(grid + inside))
        tx = [self.tax_on_gross(v) for v in pts]
        return sum((pts[i + 1] - pts[i]) * (tx[i + 1] + tx[i]) / 2 for i in range(len(pts) - 1)) / (hi - lo)

    def pareto_band_mean_tax(self, scale, alpha, plo, phi, bandmean):
        p, s = self.p, self.p["Tax_SacPct"]
        xhi = 2 * max(p["Tax_TaperZero"], p["Tax_AddThreshold"], p["Tax_NIUEL"]) / (1 - s)
        slope = (self.tax_on_gross(2 * xhi) - self.tax_on_gross(xhi)) / xhi
        resid = []
        for i in range(1, 201):
            pp = plo + (phi - plo) * (i - 0.5) / 200
            inc = scale * (0.25 / (1 - pp)) ** (1 / alpha)
            resid.append(self.tax_on_gross(inc) - slope * inc)
        return slope * bandmean + sum(resid) / len(resid)

    # ----- one pay distribution: P10..P75 plus the mean -> average tax per job -----
    @staticmethod
    def implied_alpha(pc, mean):
        b, c, d, e, f, g, h, i = pc
        avg = (0.1 * (b / 2) + 0.1 * (b + c) / 2 + 0.05 * (c + d) / 2 + 0.05 * (d + e) / 2 + 0.1 * (e + f) / 2
               + 0.2 * (f + g) / 2 + 0.1 * (g + h) / 2 + 0.05 * (h + i) / 2)
        t = mean - avg
        return t / (t - 0.25 * i)

    @staticmethod
    def tail(p75, a):
        return [p75 * (5 - 4 * (5 / 4) ** (1 / a)) / (1 - 1 / a),
                p75 * (2 * (5 / 4) ** (1 / a) - (5 / 2) ** (1 / a)) / (1 - 1 / a),
                p75 * (5 / 2) ** (1 / a) / (1 - 1 / a)]

    def band_taxes(self, pc, a, tail):
        pts = [0] + list(pc)
        out = [self.band_mean_tax(pts[k], pts[k + 1]) for k in range(8)]
        out += [self.pareto_band_mean_tax(pc[7], a, lo, hi, m) for (lo, hi), m in zip(((0.75, 0.8), (0.8, 0.9), (0.9, 1)), tail)]
        return out

    def average_tax(self, pc, mean):
        a = self.implied_alpha(pc, mean)
        tail = self.tail(pc[7], a)
        taxes = self.band_taxes(pc, a, tail)
        return sum(w * t for w, t in zip(WEIGHTS, taxes)) / 100, {"alpha": a, "tail": tail, "taxes": taxes}


def params(wb):
    ws = wb["Tax Assumptions"]
    out = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[2] and isinstance(row[1], (int, float, bool)):
            out[row[2]] = float(row[1])
    return out


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


# ----- the Wales workbook -----
def compute_wales(wb):
    """Per-constituency results for a Wales-layout workbook with no saved values (openpyxl, formulas)."""
    m = Model(params(wb))
    weeks = m.p["Annual_Weeks"] if "Annual_Weeks" in m.p else 52.14
    od, ap, jm, wl = wb["Other Data"], wb["Annual Pay Across Percentiles"], wb["Job-Mover Adjustment"], wb["Wales"]

    # Other Data: mean pay used in the model (published mean, or median x the regional mean/median ratio)
    rows = [[od.cell(r, c).value for c in range(10, 21)] for r in range(3, 35)]
    ratios = {}
    for r in rows:
        if num(r[2]) and num(r[7]):
            ratios.setdefault(r[6], []).append(r[2] / r[7])
    med = {k: statistics.median(v) for k, v in ratios.items()}
    other = {}
    for r in rows:
        name, party, mean, adults, _, _, region, median, _, code, mp = r
        used = mean if num(mean) else xround(median * med[region], 0)
        other[name.strip()] = {"party": party, "adults": adults, "mean": used, "region": region, "median": median,
                               "code": code, "mp": mp, "meanEst": not num(mean)}
    region_mean = od["B3"].value

    # Annual Pay Across Percentiles: rows by address, and the regional row
    ap_rows = {r: [ap.cell(r, c).value for c in range(1, 14)] for r in range(4, 42)}
    ap_by_name = {str(v[0]).strip(): v for r, v in ap_rows.items() if r >= 10 and v[0]}

    def ref_value(sheet, col, row):
        if sheet == "Annual Pay Across Percentiles":
            return ap.cell(row, col).value
        if sheet == "Other Data":
            if col == 15:  # mean pay used in the model
                return other[str(od.cell(row, 10).value).strip()]["mean"]
            return od.cell(row, col).value
        raise ValueError(sheet)

    def col_index(letters):
        n = 0
        for ch in letters:
            n = n * 26 + ord(ch) - 64
        return n

    ref_re = re.compile(r"(?:'([^']+)'|([A-Za-z]+))?!?\$?([A-Z]{1,2})\$?(\d+)")

    def eval_row(r):
        """Row r of the Wales sheet, B..I: each cell is a reference or simple arithmetic on the row."""
        vals = {}

        def cell(col):
            if col in vals:
                return vals[col]
            f = wl.cell(r, col).value
            if not (isinstance(f, str) and f.startswith("=")):
                vals[col] = f
                return f
            expr = f[1:]

            def sub(mo):
                sheet = mo.group(1) or mo.group(2)
                c, rr = col_index(mo.group(3)), int(mo.group(4))
                if sheet and mo.group(0).find("!") >= 0:
                    v = ref_value(sheet, c, rr)
                else:
                    v = cell(c) if rr == r else wl.cell(rr, c).value
                if not num(v):
                    raise ValueError(f"row {r}: {mo.group(0)} is not a number ({v!r})")
                return repr(float(v))

            expr = re.sub(r"(?:'[^']+'|[A-Za-z]+)!\$?[A-Z]{1,2}\$?\d+|\$?[A-Z]{1,2}\$?\d+", lambda mo: sub(ref_re.fullmatch(mo.group(0))), expr)
            v = eval(expr.replace("^", "**"), {"__builtins__": {}})
            vals[col] = v
            return v

        return [cell(c) for c in range(2, 10)]

    # job movers: extra jobs from the weekly table, taxed on the weekly distribution x 52.14
    jrows = {}
    for r in range(8, 40):
        v = [jm.cell(r, c).value for c in range(1, 17)]
        if v[0]:
            jrows[str(v[0]).strip()] = v
    jratio = statistics.median([v[6] / v[5] for v in jrows.values() if num(v[6]) and num(v[5])])

    def job_movers(name):
        v = jrows[name]
        weekly_jobs = v[2]
        annual_jobs = ap_by_name[name][1] / 1000
        extra = max(0, weekly_jobs - annual_jobs)
        mean = (v[6] if num(v[6]) else v[5] * jratio) * weeks
        wk = v[8:16]  # weekly P10, P20, P25, P30, P40, P60, P70, P75
        memo, busy = {}, set()

        def q(k):  # annualised percentile k (0..7) with the sheet's gap rules
            if k in memo:
                return memo[k]
            if num(wk[k]):
                memo[k] = wk[k] * weeks
                return memo[k]
            if k in busy:
                raise ValueError(f"{name}: circular gap-fill at weekly percentile {k}")
            busy.add(k)
            rules = {
                0: lambda: q(1) - (q(2) - q(1)) / 5 * 10,
                1: lambda: q(2) - (q(3) - q(2)) / 5 * 5,
                2: lambda: q(1) + (q(1) - q(0)) / 10 * 5,
                3: lambda: q(2) + (q(2) - q(1)) / 5 * 5,
                4: lambda: q(3) + (q(3) - q(2)) / 5 * 10,
                5: lambda: q(4) + (q(4) - q(3)) / 10 * 20,
                6: lambda: q(5) + (q(5) - q(4)) / 20 * 10,
                7: lambda: q(6) + (q(6) - q(5)) / 10 * 5,
            }
            memo[k] = rules[k]()
            busy.discard(k)
            return memo[k]

        pc = [q(k) for k in range(8)]
        avg, _ = m.average_tax(pc, mean)
        return extra * 1000, avg

    def benefit_rows(sheet):
        ws = wb[sheet]
        per = ws["I6"].value
        per = weeks if not num(per) else per
        out = {}
        for r in range(9, 41):
            b, c, d = ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value
            if b and num(c) and num(d):
                out[str(b).strip()] = (c, d * per)
        return out

    ben = [benefit_rows(s) for s in BENEFIT_SHEETS]

    seats = []
    for r in range(4, wl.max_row + 1, 22):
        name = wl.cell(r, 1).value
        if not name:
            continue
        name = str(name).strip()
        o = other[name]
        pc = eval_row(r)
        tax12, detail = m.average_tax(pc, o["mean"])
        jobs12 = ap_by_name[name][1]
        jobs_mv, tax_mv = job_movers(name) if m.p.get("Tax_IncludeJobMovers", 1) else (0, 0)
        jobs = jobs12 + jobs_mv
        avg_tax = (tax12 * jobs12 + tax_mv * jobs_mv) / jobs if jobs else tax12
        ben_amounts = [b[name][0] * b[name][1] for b in ben]
        claims = [b[name][0] for b in ben]
        imputed = sum(1 for c in range(3, 11) if ap_by_name[name][c - 1] == "x")
        seats.append({
            "name": name, "region": o["region"], "party": o["party"], "code": o["code"], "mpSheet": o["mp"],
            "tax": avg_tax * jobs, "tpw": avg_tax, "wel": sum(ben_amounts), "jobs": jobs, "adults": o["adults"],
            "ben": ben_amounts, "claims": claims, "pay": pc, "meanPay": o["mean"], "medPay": o["median"],
            "jobsSplit": [jobs12, jobs_mv], "meanEst": o["meanEst"], "imputed": imputed,
        })
    region_pay = [v for v in ap_rows[4][2:10]]
    return seats, {"regionPay": region_pay, "regionMean": region_mean}


# ----- check against the England workbook's saved results -----
def check_england(path):
    import openpyxl
    wbv = openpyxl.load_workbook(path, read_only=True, data_only=True)
    m = Model(params(wbv))
    worst = {"band": 0, "pareto": 0, "alpha": 0, "avg": 0, "jm": 0}
    n = 0
    for sheet in ["North East", "London", "South West"]:
        rows = [list(r) for r in wbv[sheet].iter_rows(values_only=True)]
        for i, r in enumerate(rows):
            if not (isinstance(r[0], str) and num(r[1]) and i + 4 < len(rows) and rows[i + 1][0] == "Mean Annual Tax in Band"):
                continue
            pc, tail_saved, alpha_saved = r[1:9], r[9:12], r[12]
            mean = rows[i + 3][1]
            saved_taxes = rows[i + 1][1:12]
            saved_c8 = rows[i + 4][2]
            a = m.implied_alpha(pc, mean)
            taxes = m.band_taxes(pc, a, m.tail(pc[7], a))
            c8 = sum(w * t for w, t in zip(WEIGHTS, taxes)) / 100
            worst["alpha"] = max(worst["alpha"], abs(a - alpha_saved) / abs(alpha_saved))
            worst["band"] = max(worst["band"], max(abs(x - y) for x, y in zip(taxes[:8], saved_taxes[:8])))
            worst["pareto"] = max(worst["pareto"], max(abs(x - y) for x, y in zip(taxes[8:], saved_taxes[8:])))
            worst["avg"] = max(worst["avg"], abs(c8 - saved_c8))
            n += 1
    jm = [list(r) for r in wbv["Job-Mover Adjustment"].iter_rows(min_row=8, values_only=True)]
    k = 0
    for r in jm:
        if not (isinstance(r[0], str) and num(r[39] if len(r) > 39 else None)):
            continue
        pc, alpha, saved_avg = r[16:24], r[27], r[39]
        if not all(num(v) for v in pc) or not num(alpha):
            continue
        taxes = m.band_taxes(pc, alpha, m.tail(pc[7], alpha))
        worst["jm"] = max(worst["jm"], abs(sum(w * t for w, t in zip(WEIGHTS, taxes)) / 100 - saved_avg))
        k += 1
    return n, k, worst


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    root = os.environ.get("WT_ROOT", os.path.join(here, ".."))
    n, k, worst = check_england(os.path.join(root, "A Tale of Two Cities (Welfare Analysis) 2026.xlsx"))
    print(f"England check: {n} constituency pay blocks and {k} job-mover rows recomputed.")
    print(f"  largest differences from the saved results: band tax £{worst['band']:.6f}, Pareto band tax £{worst['pareto']:.6f}, "
          f"implied Pareto parameter {worst['alpha']:.2e} (relative), average tax per job £{worst['avg']:.6f}, "
          f"job-mover average tax £{worst['jm']:.6f}")
